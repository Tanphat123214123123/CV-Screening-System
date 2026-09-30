"""Unit test cho cac buoc NLP con lai cua pipeline: trich xuat, cham diem, doc file."""
import pytest
from docx import Document

from app.main import parse_cv_file
from app.nlp import extractor, scorer


class TestExtractProfile:
    def test_trich_ky_nang_khong_nham_java_trong_javascript(self):
        profile = extractor.extract_profile("Skills: JavaScript, React, Node.js")
        assert "javascript" in profile["skills"]
        assert "java" not in profile["skills"]

    def test_lay_so_nam_kinh_nghiem_lon_nhat(self):
        text = "1 year at A. Total: 4+ years experience. 2 năm làm Java."
        assert extractor.extract_profile(text)["years_experience"] == 4

    def test_chan_so_nam_vo_ly(self):
        assert extractor.extract_profile("99 years of experience")["years_experience"] == 40

    def test_trich_email_va_so_dien_thoai(self):
        profile = extractor.extract_profile("Liên hệ: an.nguyen@example.com - 0912 345 678")
        assert profile["email"] == "an.nguyen@example.com"
        assert profile["phone"] is not None

    def test_cv_rong(self):
        profile = extractor.extract_profile("")
        assert profile["skills"] == []
        assert profile["years_experience"] == 0


class TestComputeScore:
    @pytest.mark.parametrize(
        ("coverage", "years", "expected"),
        [
            (1.0, 5, 100.0),   # du ky nang + du kinh nghiem = diem toi da
            (1.0, 10, 100.0),  # kinh nghiem chan tren 5 nam
            (0.5, 0, 40.0),
            (0.0, 3, 12.0),
            (0.0, -2, 0.0),    # khong am diem
        ],
    )
    def test_cong_thuc(self, coverage, years, expected):
        assert scorer.compute_score(coverage, years) == expected

    def test_tom_tat_rule_based(self):
        summary = scorer.build_summary(
            {"years_experience": 3}, {"matched": ["java"], "missing": ["aws"]}, 62.0
        )
        assert summary.startswith("Diem phu hop: 62.0/100.")
        assert "java" in summary and "aws" in summary


class TestParseCvFile:
    def test_doc_docx_ca_doan_van_va_bang(self, tmp_path):
        path = tmp_path / "cv.docx"
        doc = Document()
        doc.add_paragraph("Nguyen Van A - Java developer")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Spring Boot, Docker"
        doc.save(path)

        text = parse_cv_file(str(path), "CV.DOCX")
        assert "Java developer" in text
        assert "Spring Boot, Docker" in text

    @pytest.mark.parametrize("name", ["cv.txt", "cv", "cv.exe"])
    def test_tu_choi_dinh_dang_khong_ho_tro(self, tmp_path, name):
        with pytest.raises(ValueError):
            parse_cv_file(str(tmp_path / name), name)
