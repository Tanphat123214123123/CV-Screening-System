"""Trich xuat text tu file Word (.docx) bang python-docx."""
from docx import Document


def extract_text(path: str) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Doc them noi dung trong bang (nhieu CV trinh bay bang table)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)
